"""Lectura de los extractos (hojas de cálculo que mantiene Administración en Drive).

Banco (una pestaña, p. ej. SABADELL): columnas "F. Operativa", "Concepto", "F. Valor",
"Importe" (negativo = cargo), "Saldo" y, después, notas y referencias. Las fechas vienen como
texto "dd/mm/aaaa" o, si Excel las convirtió, como fecha con el día y el mes intercambiados
(09/10/2026 → 10 de septiembre): se deshace el cambio.

Tarjeta (un fichero por año, una pestaña por mes: ENE…DIC): "Fecha" ("dd/mm", o fecha con el
día y el mes intercambiados), "Concepto", población, "Importe" como texto ("1.010,05") y
moneda (en algunas pestañas con una columna vacía en medio). Los cargos en otra moneda llevan
debajo una fila "Importe original: 49,00 USD(…)".
"""

from __future__ import annotations

import hashlib
import io
import re
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal

from ..extraccion.base import ExtraccionError
from ..extraccion.base import importe as leer_importe
from ..models import MetodoPago, MovimientoPago

MESES = ["ENE", "FEB", "MAR", "ABR", "MAY", "JUN", "JUL", "AGO", "SEP", "OCT", "NOV", "DIC"]
ORIGINAL = re.compile(r"Importe original:\s*(?P<importe>[\d.,]+)\s*(?P<moneda>[A-Z]{3})")
# Liquidación mensual de la tarjeta en el banco: sus cargos se concilian con la hoja de la tarjeta
LIQUIDACION_TARJETA = re.compile(r"^TARJETA CREDITO", re.I)


class HojaError(Exception):
    """La hoja no tiene el formato esperado."""


@dataclass
class Movimiento:
    origen: str
    fecha: date
    concepto: str
    importe: Decimal  # positivo = pago
    metodo: str
    hoja: str
    importe_original: Decimal | None = None
    moneda_original: str = ""
    clave: str = field(default="", compare=False)


def _texto(valor) -> str:
    return "" if valor is None else str(valor).strip()


def _fecha_banco(valor) -> date:
    if isinstance(valor, datetime):
        return date(valor.year, valor.day, valor.month)  # Excel cambió día y mes
    if m := re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{4})", _texto(valor)):
        return date(int(m[3]), int(m[2]), int(m[1]))
    raise HojaError(f"Fecha que no entiendo: {valor!r}")


def _dia_tarjeta(valor, mes: int) -> tuple[int, int] | None:
    """(día, mes) de la fecha de un cargo de tarjeta, o None si la fila no tiene fecha."""
    if valor is None or valor == "":
        return None
    if isinstance(valor, datetime):
        # "05/03" en la pestaña MAR se convirtió en el 3 de mayo (o se quedó bien)
        if valor.day == mes:
            return valor.month, mes
        return valor.day, valor.month
    if isinstance(valor, int | float):
        return int(valor), mes
    if m := re.fullmatch(r"(\d{1,2})/(\d{1,2})(?:/\d{2,4})?", _texto(valor)):
        return int(m[1]), int(m[2])
    raise HojaError(f"Fecha que no entiendo: {valor!r}")


def _importe_tarjeta(fila: tuple):
    """El importe va justo antes de la moneda ("EUR"); según la pestaña, en la columna D o E."""
    for i, valor in enumerate(fila):
        if i > 1 and re.fullmatch(r"[A-Z]{3}", _texto(valor)):
            anteriores = [v for v in fila[2:i] if v is not None]
            return anteriores[-1] if anteriores else None
    return fila[4] if len(fila) > 4 else None


def metodo_de_banco(concepto: str) -> str:
    if re.match(r"(TRANSFERENCIA|TRANSF\.|TRF\.)", concepto, re.I):
        return MetodoPago.TRANSFERENCIA
    return MetodoPago.DOMICILIACION


def leer_banco(contenido: bytes, pestana: str, nombre: str = "banco") -> list[Movimiento]:
    import openpyxl

    libro = openpyxl.load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
    if pestana not in libro.sheetnames:
        raise HojaError(f"No encuentro la pestaña «{pestana}» en {nombre}")
    movimientos = []
    for n, fila in enumerate(libro[pestana].iter_rows(values_only=True), start=1):
        if n == 1 or not fila or fila[0] in (None, "") or fila[3] in (None, ""):
            continue
        concepto = _texto(fila[1])
        if LIQUIDACION_TARJETA.match(concepto):
            continue
        referencias = " ".join(_texto(v) for v in fila[7:9] if _texto(v))
        movimientos.append(
            Movimiento(
                origen=MovimientoPago.Origen.BANCO,
                fecha=_fecha_banco(fila[0]),
                concepto=f"{concepto} {referencias}".strip()[:500],
                importe=-Decimal(str(fila[3])).quantize(Decimal("0.01")),
                metodo=metodo_de_banco(concepto),
                hoja=f"{nombre} › {pestana} › fila {n}",
            )
        )
    return movimientos


def leer_tarjeta(
    contenido: bytes, anio: int, nombre: str = "tarjeta", avisos: list[str] | None = None
) -> list[Movimiento]:
    """Cargos de la tarjeta. Las filas que no se entienden se saltan y se anotan en `avisos`."""
    import openpyxl

    libro = openpyxl.load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
    movimientos: list[Movimiento] = []
    for numero, pestana in enumerate(MESES, start=1):
        if pestana not in libro.sheetnames:
            continue
        for n, fila in enumerate(libro[pestana].iter_rows(values_only=True), start=1):
            fila = tuple(None if _texto(v) == "" else v for v in fila)
            if n == 1 or not any(v is not None for v in fila):
                continue
            concepto = _texto(fila[1])
            if (m := ORIGINAL.search(concepto)) and movimientos and fila[0] in (None, ""):
                movimientos[-1].importe_original = leer_importe(m["importe"], decimal=",")
                movimientos[-1].moneda_original = m["moneda"]
                continue
            try:
                dia_mes = _dia_tarjeta(fila[0], numero)
                bruto = _importe_tarjeta(fila)
                if dia_mes is None or bruto is None:
                    continue
                valor = (
                    Decimal(str(bruto))
                    if isinstance(bruto, int | float)
                    else leer_importe(str(bruto), decimal=",")
                )
                dia, mes = dia_mes
                date(anio, mes, dia)  # valida la fecha
            except (ExtraccionError, HojaError, ValueError) as exc:
                if avisos is not None:
                    avisos.append(f"{nombre} › {pestana} › fila {n}: {exc}")
                continue
            # Un cargo de diciembre en la pestaña de enero es del año anterior
            anio_cargo = anio - 1 if mes > numero + 1 else anio
            movimientos.append(
                Movimiento(
                    origen=MovimientoPago.Origen.TARJETA,
                    fecha=date(anio_cargo, mes, dia),
                    concepto=concepto[:500],
                    importe=valor.quantize(Decimal("0.01")),
                    metodo=MetodoPago.TARJETA,
                    hoja=f"{nombre} › {pestana} › fila {n}",
                )
            )
    return movimientos


def asignar_claves(movimientos: list[Movimiento]) -> list[Movimiento]:
    """Clave estable de cada movimiento: sus datos más su nº de orden entre los idénticos."""
    vistos: Counter = Counter()
    for m in movimientos:
        base = f"{m.origen}|{m.fecha.isoformat()}|{m.concepto}|{m.importe}"
        vistos[base] += 1
        m.clave = hashlib.sha1(f"{base}|{vistos[base]}".encode()).hexdigest()
    return movimientos
